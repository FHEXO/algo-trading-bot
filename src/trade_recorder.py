import os
from datetime import datetime
from typing import Any, Dict, List, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class TradeRecorder:
    """Registra, formatea y exporta todas las operaciones ejecutadas por el bot a archivos CSV y Excel (.xlsx)."""

    COLUMNS: List[str] = [
        "ID",
        "Fecha_Entrada",
        "Fecha_Salida",
        "Simbolo",
        "Tipo",
        "Cantidad",
        "Precio_Entrada_USD",
        "Precio_Salida_USD",
        "Stop_Loss_USD",
        "Take_Profit_USD",
        "Motivo_Salida",
        "PnL_USD",
        "PnL_Porcentaje",
        "Balance_Total_USD",
        "Modo_Proteccion",
    ]

    def __init__(
        self,
        csv_file: str = "data/registro_operaciones.csv",
        excel_file: str = "data/registro_operaciones.xlsx",
    ) -> None:
        self.csv_file = csv_file
        self.excel_file = excel_file

        os.makedirs(os.path.dirname(self.csv_file), exist_ok=True)
        self._init_files()

    def _init_files(self) -> None:
        """Inicializa los archivos con sus encabezados si no existen."""
        if not os.path.exists(self.csv_file):
            df_empty = pd.DataFrame(columns=self.COLUMNS)
            df_empty.to_csv(self.csv_file, index=False, encoding="utf-8-sig")

        if not os.path.exists(self.excel_file):
            try:
                df_empty = pd.DataFrame(columns=self.COLUMNS)
                df_empty.to_excel(self.excel_file, index=False)
                self._apply_excel_styles()
            except Exception as e:
                print(f"[WARN TradeRecorder] No se pudo inicializar Excel: {e}")

    def get_trades_df(self) -> pd.DataFrame:
        """Carga el DataFrame de operaciones registradas."""
        if os.path.exists(self.csv_file):
            try:
                return pd.read_csv(self.csv_file, encoding="utf-8-sig")
            except Exception:
                return pd.DataFrame(columns=self.COLUMNS)
        return pd.DataFrame(columns=self.COLUMNS)

    def _apply_excel_styles(self) -> None:
        """Aplica un diseño profesional institucinal al archivo Excel (.xlsx)."""
        if not os.path.exists(self.excel_file):
            return

        try:
            wb = openpyxl.load_workbook(self.excel_file)
            ws = wb.active

            # Paleta de Estilos Profesionales
            header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            data_font = Font(name="Segoe UI", size=10)
            bold_data_font = Font(name="Segoe UI", size=10, bold=True)

            green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
            green_font = Font(name="Segoe UI", size=10, bold=True, color="155724")

            red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
            red_font = Font(name="Segoe UI", size=10, bold=True, color="721C24")

            align_center = Alignment(horizontal="center", vertical="center")
            align_right = Alignment(horizontal="right", vertical="center")

            thin_border = Border(
                left=Side(style="thin", color="D3D3D3"),
                right=Side(style="thin", color="D3D3D3"),
                top=Side(style="thin", color="D3D3D3"),
                bottom=Side(style="thin", color="D3D3D3"),
            )

            # Estilizar Encabezados (Fila 1)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center

            # Estilizar Filas de Datos
            pnl_col_idx = self.COLUMNS.index("PnL_USD") + 1 if "PnL_USD" in self.COLUMNS else 12

            for row_idx in range(2, ws.max_row + 1):
                pnl_val_cell = ws.cell(row=row_idx, column=pnl_col_idx)
                pnl_str = str(pnl_val_cell.value or "")
                is_positive = pnl_str.startswith("+")
                is_negative = pnl_str.startswith("-")

                for col_idx in range(1, len(self.COLUMNS) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    col_name = self.COLUMNS[col_idx - 1] if col_idx - 1 < len(self.COLUMNS) else ""

                    cell.font = data_font
                    cell.border = thin_border

                    # Alineaciones según tipo de campo
                    if col_name in ["ID", "Fecha_Entrada", "Fecha_Salida", "Simbolo", "Tipo", "Motivo_Salida", "Modo_Proteccion"]:
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_right

                    # Formato condicional para PnL ($ y %)
                    if col_name in ["PnL_USD", "PnL_Porcentaje"]:
                        if is_positive:
                            cell.fill = green_fill
                            cell.font = green_font
                        elif is_negative:
                            cell.fill = red_fill
                            cell.font = red_font
                        else:
                            cell.font = bold_data_font

            # Auto-ajuste dinámico de ancho de columnas
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(self.excel_file)
        except Exception as e:
            print(f"[WARN TradeRecorder] Error al aplicar estilos a Excel: {e}")

    def record_trade(
        self,
        entry_time: str,
        exit_time: str,
        symbol: str,
        amount: float,
        entry_price: float,
        exit_price: float,
        stop_loss_price: float,
        take_profit_price: float,
        exit_reason: str,
        balance_total: float,
        protection_mode: str = "NONE",
        trade_type: str = "BUY",
    ) -> Dict[str, Any]:
        """Calcula el PnL, registra la operación en CSV/Excel y aplica diseño visual."""
        df_current = self.get_trades_df()
        next_id = len(df_current) + 1

        # Cálculo de métricas
        pnl_usd = round((exit_price - entry_price) * amount, 4)
        pnl_pct = round(((exit_price - entry_price) / entry_price) * 100, 2)

        trade_row: Dict[str, Any] = {
            "ID": next_id,
            "Fecha_Entrada": entry_time,
            "Fecha_Salida": exit_time,
            "Simbolo": symbol,
            "Tipo": trade_type,
            "Cantidad": round(amount, 6),
            "Precio_Entrada_USD": round(entry_price, 2),
            "Precio_Salida_USD": round(exit_price, 2),
            "Stop_Loss_USD": round(stop_loss_price, 2),
            "Take_Profit_USD": round(take_profit_price, 2),
            "Motivo_Salida": exit_reason,
            "PnL_USD": f"{'+' if pnl_usd >= 0 else ''}{pnl_usd:.2f}",
            "PnL_Porcentaje": f"{'+' if pnl_pct >= 0 else ''}{pnl_pct:.2f}%",
            "Balance_Total_USD": round(balance_total, 2),
            "Modo_Proteccion": protection_mode,
        }

        # Actualizar CSV
        df_new = pd.DataFrame([trade_row])
        df_updated = pd.concat([df_current, df_new], ignore_index=True)
        df_updated.to_csv(self.csv_file, index=False, encoding="utf-8-sig")

        # Actualizar Excel (.xlsx) y aplicar formato visual
        try:
            df_updated.to_excel(self.excel_file, index=False)
            self._apply_excel_styles()
        except Exception as e:
            print(f"[WARN TradeRecorder] Error al guardar .xlsx: {e}")

        return trade_row


if __name__ == "__main__":
    recorder = TradeRecorder()
    print("TradeRecorder inicializado correctamente con formateador visual.")
    print(f"Archivo CSV:   {recorder.csv_file}")
    print(f"Archivo Excel: {recorder.excel_file}")
